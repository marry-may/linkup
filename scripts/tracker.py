#!/usr/bin/env python3
"""Manage docs/project_tracker.xlsx safely.

Examples:
    python scripts/tracker.py next
    python scripts/tracker.py list --status "Not Started" [--phase P2] [--priority P0]
    python scripts/tracker.py show P1-T02
    python scripts/tracker.py update P1-T03 --status Done --note "tests passing"
    python scripts/tracker.py decision "Use X" --options "X, Y" --rationale "Because..."
"""
import argparse
import datetime as dt
import os
import sys
import tempfile
import textwrap
from pathlib import Path

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Border, Side
except ImportError:
    sys.exit("openpyxl missing. Run: python3 -m venv .venv && .venv/bin/pip install -r scripts/requirements.txt")

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "docs" / "project_tracker.xlsx"
STATUSES = ["Not Started", "In Progress", "Blocked", "In Review", "Done"]
OPEN = {"Not Started", "In Progress", "Blocked", "In Review"}
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(bottom=Side(style="thin", color="E5E7EB"))


class Tracker:
    def __init__(self, path=XLSX):
        if not path.exists():
            sys.exit("%s not found. Run: python scripts/build_tracker.py" % path.relative_to(ROOT))
        self.path = path
        self.wb = load_workbook(path)
        self.ws = self.wb["Tasks"]
        self.cols = {c.value: c.column for c in self.ws[1]}
        self.tasks = []  # list of (row_number, dict)
        for row in self.ws.iter_rows(min_row=2):
            vals = {h: row[i - 1].value for h, i in self.cols.items()}
            if vals.get("Task ID"):
                self.tasks.append((row[0].row, vals))
        self.by_id = {t["Task ID"]: (r, t) for r, t in self.tasks}
        self.tasks_by_row = {r: t for r, t in self.tasks}

    # ---- helpers -------------------------------------------------------
    @staticmethod
    def deps(task):
        return [d.strip() for d in str(task.get("Depends On") or "").split(",") if d.strip()]

    def unmet(self, task):
        return [d for d in self.deps(task) if d not in self.by_id or self.by_id[d][1]["Status"] != "Done"]

    def set(self, row, header, value, date=False):
        c = self.ws.cell(row=row, column=self.cols[header])
        c.value = value  # ws.cell(value=None) would not clear the cell
        c.alignment = WRAP
        self.tasks_by_row[row][header] = value
        if date:
            c.number_format = "yyyy-mm-dd"

    def append_row(self, sheet, values, date_col=1):
        ws = self.wb[sheet]
        r = ws.max_row + 1
        while r > 2 and all(c.value in (None, "") for c in ws[r - 1]):
            r -= 1
        for i, v in enumerate(values, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.alignment, c.border = WRAP, THIN
            if i == date_col:
                c.number_format = "yyyy-mm-dd"
        if ws.auto_filter.ref:
            start = ws.auto_filter.ref.split(":")[0]
            end_col = ws.auto_filter.ref.split(":")[1].rstrip("0123456789")
            ws.auto_filter.ref = "%s:%s%d" % (start, end_col, r)

    def touch_overview(self):
        ov = self.wb["Overview"]
        for row in ov.iter_rows():
            if row[0].value == "Last updated":
                row[1].value = dt.date.today()
                row[1].number_format = "yyyy-mm-dd"

    def save(self):
        self.touch_overview()
        fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=str(self.path.parent))
        os.close(fd)
        try:
            self.wb.save(tmp)
            os.replace(tmp, self.path)
        except PermissionError:
            os.unlink(tmp)
            sys.exit("Could not write tracker — close it in Excel/Numbers and retry.")

    def progress(self):
        mvp = [t for _, t in self.tasks if t["Phase"] != "P6"]
        done = sum(1 for t in mvp if t["Status"] == "Done")
        return done, len(mvp)

    def current_phase(self):
        for _, t in self.tasks:
            if t["Status"] != "Done":
                return t["Phase"]
        return None


# ---- commands --------------------------------------------------------------
def fmt_row(t):
    sec = " [SEC]" if t.get("Security Related (Y/N)") == "Y" else ""
    return "%-7s %-3s %-12s %s%s" % (t["Task ID"], t["Priority"], t["Status"], t["Title"], sec)


def cmd_list(tr, args):
    rows = [t for _, t in tr.tasks
            if (not args.status or t["Status"] == args.status)
            and (not args.phase or t["Phase"] == args.phase)
            and (not args.priority or t["Priority"] == args.priority)]
    for t in rows:
        print(fmt_row(t))
    print("\n%d task(s)" % len(rows))


def cmd_show(tr, args):
    if args.task_id not in tr.by_id:
        sys.exit("Unknown task: %s" % args.task_id)
    t = tr.by_id[args.task_id][1]
    for h in ["Task ID", "Phase", "Title", "Status", "Priority", "Depends On", "Security Related (Y/N)", "OWASP Ref",
              "Est. Effort", "Description", "Acceptance Criteria", "Files Touched", "Notes", "Date Completed"]:
        v = t.get(h)
        if isinstance(v, dt.datetime):
            v = v.date()
        if v in (None, ""):
            continue
        label = h + ":"
        if len(str(v)) > 70:
            print(label)
            print(textwrap.indent(textwrap.fill(str(v), 100), "    "))
        else:
            print("%-24s %s" % (label, v))
    unmet = tr.unmet(t)
    if unmet and t["Status"] != "Done":
        print("\n⚠ Unmet dependencies: %s" % ", ".join(unmet))


def cmd_next(tr, args):
    done, total = tr.progress()
    print("LinkUp MVP progress: %d/%d tasks done (%d%%)\n" % (done, total, round(100 * done / total) if total else 0))
    in_progress = [t for _, t in tr.tasks if t["Status"] in ("In Progress", "In Review")]
    if in_progress:
        print("Already in progress — finish these first:")
        for t in in_progress:
            print("  " + fmt_row(t))
        print()
    ready = [t for _, t in tr.tasks if t["Status"] == "Not Started" and not tr.unmet(t)]
    if not ready:
        blocked = [t for _, t in tr.tasks if t["Status"] in OPEN]
        print("No ready tasks." if blocked else "All tasks are Done 🎉")
        return
    # Prefer lower phase, then priority, then ID order (tracker order)
    order = {tid: i for i, (_, t) in enumerate(tr.tasks) for tid in [t["Task ID"]]}
    ready.sort(key=lambda t: (t["Phase"], t["Priority"], order[t["Task ID"]]))
    nxt = ready[0]
    print("NEXT TASK")
    print("=" * 60)
    args.task_id = nxt["Task ID"]
    cmd_show(tr, args)
    others = [t["Task ID"] for t in ready[1:6]]
    if others:
        print("\nAlso ready: " + ", ".join(others))


def cmd_update(tr, args):
    if args.task_id not in tr.by_id:
        sys.exit("Unknown task: %s" % args.task_id)
    if not args.status and not args.note:
        sys.exit("Nothing to update: pass --status and/or --note")
    row, t = tr.by_id[args.task_id]
    today = dt.date.today()
    old = t["Status"]
    if args.status:
        if args.status not in STATUSES:
            sys.exit("Invalid status. Choose one of: %s" % ", ".join(STATUSES))
        unmet = tr.unmet(t)
        if args.status in ("In Progress", "In Review", "Done") and unmet and not args.force:
            sys.exit("Refusing: %s has unmet dependencies (%s). Use --force to override." % (args.task_id, ", ".join(unmet)))
        tr.set(row, "Status", args.status)
        if args.status == "Done" and old != "Done":
            tr.set(row, "Date Completed", today, date=True)
            summary = "Marked Done" + (": " + args.note if args.note else "")
            tr.append_row("Changelog", [today, args.task_id, summary])
        elif args.status != "Done" and old == "Done":
            tr.set(row, "Date Completed", None)
            tr.append_row("Changelog", [today, args.task_id, "Reopened (%s → %s)" % (old, args.status)])
    if args.note:
        prev = t.get("Notes") or ""
        entry = "[%s] %s" % (today.isoformat(), args.note)
        tr.set(row, "Notes", (prev + "\n" + entry).strip())
    tr.save()
    print("%s: %s → %s%s" % (args.task_id, old, args.status or old, " (note added)" if args.note else ""))
    if args.status == "Done":
        done, total = tr.progress()
        print("MVP progress: %d/%d done. Changelog entry added." % (done, total))


def cmd_decision(tr, args):
    tr.append_row("Decisions", [dt.date.today(), args.decision, args.options or "", args.rationale or ""])
    tr.append_row("Changelog", [dt.date.today(), "—", "Decision recorded: " + args.decision])
    tr.save()
    print("Decision recorded.")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("list", help="list tasks")
    p.add_argument("--status", choices=STATUSES)
    p.add_argument("--phase")
    p.add_argument("--priority", choices=["P0", "P1", "P2"])
    p.set_defaults(fn=cmd_list)

    p = sub.add_parser("next", help="show the next ready task (dependencies Done)")
    p.set_defaults(fn=cmd_next)

    p = sub.add_parser("show", help="show one task in full")
    p.add_argument("task_id")
    p.set_defaults(fn=cmd_show)

    p = sub.add_parser("update", help="update a task's status and/or add a note")
    p.add_argument("task_id")
    p.add_argument("--status", choices=STATUSES)
    p.add_argument("--note")
    p.add_argument("--force", action="store_true", help="allow starting/completing with unmet dependencies")
    p.set_defaults(fn=cmd_update)

    p = sub.add_parser("decision", help="record an architecture decision")
    p.add_argument("decision")
    p.add_argument("--options")
    p.add_argument("--rationale")
    p.set_defaults(fn=cmd_decision)

    args = ap.parse_args()
    args.fn(Tracker(), args)


if __name__ == "__main__":
    main()
